# -*- coding: utf-8 -*-
"""
Lab 8-3 — Non-linear extraction: the slide and the presenter notes (Sophie)

Learning objective
------------------
A slide deck is a NON-LINEAR and elliptical source: the bullets say little, and
the real knowledge often sits in the presenter notes. Here both the visible text
and the hidden notes are extracted, and associated with the same fragment.

No API key. Dependency: python-pptx.
Run generate_sample_docs.py first.
"""

import json
from pathlib import Path
from typing import Dict, List

DOCS = Path(__file__).resolve().parent / "sample_docs"
PPTX = DOCS / "sophie_remote_work_agreement.pptx"


def extract_presentation(path: Path) -> Dict:
    from pptx import Presentation

    prs = Presentation(str(path))
    pivot = {"title": "Remote work agreement (slide deck)",
             "sections": [],
             "metadata": {"source_format": "pptx"}}

    for index, slide in enumerate(prs.slides, start=1):
        title = ""
        bullets: List[str] = []

        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            lines = [l.strip() for l in text.splitlines() if l.strip()]
            # The heuristic: the first short unlisted block is the title.
            if not title and len(lines) == 1 and not lines[0].startswith("-"):
                title = lines[0]
            else:
                for l in lines:
                    bullets.append(l.lstrip("- ").strip())

        # The presenter notes: the hidden text, often the richer one.
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        pivot["sections"].append({
            "title": title or f"Diapositive {index}",
            "position": index,
            "bullets": bullets,
            "presenter_notes": notes,
            # The usable text combines bullets AND notes.
            "text": " ".join(bullets + ([notes] if notes else [])),
        })

    return pivot


def main() -> None:
    print("=" * 78)
    print("Lab 8-3 — The slide and the presenter notes (Sophie)")
    print("=" * 78)

    if not PPTX.exists():
        print("\nDocument introuvable. Lancez d'abord : python generate_sample_docs.py")
        return

    pivot = extract_presentation(PPTX)
    print(f"\nDiapositives extraites : {len(pivot['sections'])}")

    for s in pivot["sections"]:
        print(f"\n--- Diapositive {s['position']} : {s['title']} ---")
        print("  Puces :")
        for p in s["bullets"]:
            print(f"    - {p}")
        if s["presenter_notes"]:
            print("  Presenter notes (the hidden text):")
            print(f"    {s['presenter_notes'][:160]}…")

    # The demonstration: what the bullets alone do NOT say.
    print("\n" + "=" * 78)
    print("WHY THE NOTES MATTER")
    print("=" * 78)
    bullets_only = " ".join(p for s in pivot["sections"] for p in s["bullets"]).lower()
    with_notes = " ".join(s["text"] for s in pivot["sections"]).lower()
    print("Question: \"Is the number of days settled in the deck?\"")
    in_bullets = "approved" in bullets_only or "still being" in bullets_only
    in_notes = "still being approved" in with_notes or "do not announce" in with_notes
    print(f"- Answer findable in the BULLETS alone: {'yes' if in_bullets else 'NO'}")
    print(f"- Answer findable with the NOTES: {'yes' if in_notes else 'no'}")

    out_path = DOCS / "pivot_pptx.json"
    out_path.write_text(json.dumps(pivot, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\nJSON written: {out_path.name}")

    # --- Of the pivot to the chunks (the notes are incluses in the content) ---
    from chunking import chunk_pivot, preview
    chunks = chunk_pivot(pivot, prefix="CHUNK_PPTX")
    print(f"\n--- Chunks produits : {len(chunks)} ---")
    preview(chunks)
    chunks_path = DOCS / "chunks_pptx.json"
    chunks_path.write_text(json.dumps(chunks, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Chunks written: {chunks_path.name}")

    print("\nA RETENIR")
    print("- A deck is elliptical: the bullets summarise, the notes make it precise.")
    print("- Ignoring the notes means losing half the organisation's knowledge.")
    print("- Each slide chunk folds in bullets AND notes, so it stays usable.")


if __name__ == "__main__":
    main()
